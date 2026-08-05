#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = [
#     "openpyxl>=3.1",
# ]
# ///
"""Generate the data-extraction Excel workbook for the AI/ML-PROMs-after-TJA review.

Builds /extraction/data_extraction.xlsx with two sheets:

    1. "Data Extraction" — one row per included study, with dropdown
       (list) validation on the Joint, External Validation, and PROBAST
       columns so reviewers can't free-type an off-list value.
    2. "PROBAST Guide" — the four PROBAST domains and their signaling
       questions, as a standing reference while extracting/rating.

Per project policy this script is built but is NOT auto-run — the
sheet is only generated when the user explicitly asks for it, so a
placeholder workbook doesn't get committed and then immediately
overwritten mid-extraction.

Usage:
    uv run extraction_template.py
    uv run extraction_template.py --output data_extraction.xlsx --rows 100
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

EXTRACTION_COLUMNS = [
    "Author",
    "Year",
    "Journal",
    "Country",
    "Study Design",
    "Sample Size",
    "Joint (TKA/THA/Both)",
    "Indication",
    "ML Algorithm",
    "Input Variables",
    "PROM Instrument",
    "Outcome Type (MCID vs Raw Score)",
    "Follow-up Timepoint",
    "AUC",
    "Sensitivity",
    "Specificity",
    "Accuracy",
    "R²",
    "External Validation (Yes/No)",
    "Training/Test Split",
    "Top Predictive Variables",
    "PROBAST Domain 1",
    "PROBAST Domain 2",
    "PROBAST Domain 3",
    "PROBAST Domain 4",
    "Overall PROBAST Risk",
    "Notes",
]

JOINT_OPTIONS = ["TKA", "THA", "Both"]
EXTERNAL_VALIDATION_OPTIONS = ["Yes", "No", "Not Reported"]
PROBAST_RATING_OPTIONS = ["Low", "High", "Unclear"]

DEFAULT_DATA_ROWS = 60

PROBAST_GUIDE = [
    (
        "Domain 1: Participants",
        "1.1 Were appropriate data sources used (e.g. cohort, RCT, or nested case-control study)?",
        "Registry or single/multi-center cohort with prospective PROM collection = Low.",
    ),
    (
        "Domain 1: Participants",
        "1.2 Were all inclusions and exclusions of participants appropriate?",
        "Check for selective exclusion of missing-data or poor-outcome patients that could bias the model.",
    ),
    (
        "Domain 2: Predictors",
        "2.1 Were predictors defined and assessed in a similar way for all participants?",
        "Same instruments/timing across the whole cohort = Low.",
    ),
    (
        "Domain 2: Predictors",
        "2.2 Were predictor assessments made without knowledge of outcome data?",
        "Predictors collected pre-op/at baseline, before the PROM outcome is known = Low.",
    ),
    (
        "Domain 2: Predictors",
        "2.3 Are all predictors available at the time the model is intended to be used?",
        "If the model is meant for pre-op counseling, any predictor only known post-op = High.",
    ),
    (
        "Domain 3: Outcome",
        "3.1 Was the outcome determined appropriately?",
        "Validated PROM (KOOS/HOOS/WOMAC/OHS/OKS) collected per standard protocol = Low.",
    ),
    (
        "Domain 3: Outcome",
        "3.2 Was a prespecified or standard outcome definition used?",
        "A published MCID threshold or the raw validated instrument score = Low.",
    ),
    (
        "Domain 3: Outcome",
        "3.3 Were predictors excluded from the outcome definition?",
        "Outcome definition must not double-count a predictor variable.",
    ),
    (
        "Domain 3: Outcome",
        "3.4 Was the outcome defined and determined in a similar way for all participants?",
        "Same follow-up instrument/timepoint applied cohort-wide = Low.",
    ),
    (
        "Domain 3: Outcome",
        "3.5 Was the outcome determined without knowledge of predictor information?",
        "Usually Low for PROM outcomes (patient self-report, blind to model inputs).",
    ),
    (
        "Domain 3: Outcome",
        "3.6 Was the time interval between predictor assessment and outcome determination appropriate?",
        "Follow-up timepoint clinically meaningful (e.g. 1yr) and consistent = Low.",
    ),
    (
        "Domain 4: Analysis",
        "4.1 Were there a reasonable number of participants with the outcome?",
        "Check events-per-variable / class balance for the ML task; small-N = High risk.",
    ),
    (
        "Domain 4: Analysis",
        "4.2 Were continuous and categorical predictors handled appropriately?",
        "Arbitrary dichotomization of continuous predictors without justification = High.",
    ),
    (
        "Domain 4: Analysis",
        "4.3 Were all enrolled participants included in the analysis?",
        "Large unexplained attrition between enrollment and modeled sample = High.",
    ),
    (
        "Domain 4: Analysis",
        "4.4 Were participants with missing data handled appropriately?",
        "Multiple imputation or documented complete-case rationale = Low; silent dropping = High.",
    ),
    (
        "Domain 4: Analysis",
        "4.5 Was selection of predictors based on univariable analysis avoided?",
        "Univariable pre-screening of predictors before multivariable/ML modeling = High.",
    ),
    (
        "Domain 4: Analysis",
        "4.6 Were complexities in the data accounted for appropriately?",
        "Censoring, competing risks, or non-random sampling of controls, if present, handled correctly.",
    ),
    (
        "Domain 4: Analysis",
        "4.7 Were relevant model performance measures evaluated appropriately?",
        "Discrimination (AUC) AND calibration reported, ideally on a held-out/test set = Low.",
    ),
    (
        "Domain 4: Analysis",
        "4.8 Were model overfitting, underfitting, and optimism accounted for?",
        "Cross-validation, bootstrap optimism correction, or external validation = Low; none = High.",
    ),
    (
        "Domain 4: Analysis",
        "4.9 Do predictors and their weights correspond to the multivariable analysis results?",
        "Reported final model matches the analysis actually performed (no undisclosed post-hoc tuning).",
    ),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")


def _autofit_columns(ws, headers: list[str], min_width: int = 12, max_width: int = 40) -> None:
    for idx, header in enumerate(headers, start=1):
        width = max(min_width, min(max_width, len(header) + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_data_extraction_sheet(wb: Workbook, num_rows: int) -> None:
    ws = wb.active
    ws.title = "Data Extraction"

    for col_idx, header in enumerate(EXTRACTION_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.freeze_panes = "A2"
    _autofit_columns(ws, EXTRACTION_COLUMNS)
    ws.row_dimensions[1].height = 32

    last_row = num_rows + 1  # header is row 1

    def add_dropdown(column_name: str, options: list[str]) -> None:
        col_idx = EXTRACTION_COLUMNS.index(column_name) + 1
        col_letter = get_column_letter(col_idx)
        dv = DataValidation(
            type="list",
            formula1=f'"{",".join(options)}"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Invalid entry",
            error=f"Choose one of: {', '.join(options)}",
        )
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{last_row}")

    add_dropdown("Joint (TKA/THA/Both)", JOINT_OPTIONS)
    add_dropdown("External Validation (Yes/No)", EXTERNAL_VALIDATION_OPTIONS)
    add_dropdown("PROBAST Domain 1", PROBAST_RATING_OPTIONS)
    add_dropdown("PROBAST Domain 2", PROBAST_RATING_OPTIONS)
    add_dropdown("PROBAST Domain 3", PROBAST_RATING_OPTIONS)
    add_dropdown("PROBAST Domain 4", PROBAST_RATING_OPTIONS)
    add_dropdown("Overall PROBAST Risk", PROBAST_RATING_OPTIONS)

    for row in range(2, last_row + 1):
        for col in range(1, len(EXTRACTION_COLUMNS) + 1):
            ws.cell(row=row, column=col).alignment = WRAP_ALIGN


def build_probast_guide_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("PROBAST Guide")

    headers = ["Domain", "Signaling Question", "Rating Guidance"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.freeze_panes = "A2"

    for row_idx, (domain, question, guidance) in enumerate(PROBAST_GUIDE, start=2):
        ws.cell(row=row_idx, column=1, value=domain).alignment = WRAP_ALIGN
        ws.cell(row=row_idx, column=2, value=question).alignment = WRAP_ALIGN
        ws.cell(row=row_idx, column=3, value=guidance).alignment = WRAP_ALIGN

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60

    note_row = len(PROBAST_GUIDE) + 3
    ws.cell(
        row=note_row,
        column=1,
        value=(
            "Overall PROBAST risk of bias = High if ANY domain is High. "
            "Low only if ALL four domains are Low. Otherwise Unclear."
        ),
    ).font = Font(italic=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="data_extraction.xlsx")
    parser.add_argument("--rows", type=int, default=DEFAULT_DATA_ROWS)
    args = parser.parse_args()

    output_path = Path(__file__).resolve().parent / args.output

    wb = Workbook()
    build_data_extraction_sheet(wb, args.rows)
    build_probast_guide_sheet(wb)
    wb.save(output_path)

    print(f"Wrote {output_path}")
    print(f"Data Extraction sheet: {len(EXTRACTION_COLUMNS)} columns, {args.rows} blank data rows.")
    print(f"PROBAST Guide sheet: {len(PROBAST_GUIDE)} signaling questions across 4 domains.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
